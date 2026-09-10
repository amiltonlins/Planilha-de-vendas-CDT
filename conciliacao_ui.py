"""Conciliação screens. Google Sheets is only read; no credential is exported."""
import html
import io
import threading
import time
from calendar import monthrange
from datetime import date, datetime
from decimal import Decimal

from conciliacao_visual import integer, STYLE, daily_color, projection_color, ranking_html, ranking_png, regimes_html, summary_html

from conciliacao import TZ, aggregate, award_for, normalized, read_source, summarize, useful_days, weeks


def money(value):
    return "R$ " + f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


class SourceCache:
    """Keep last complete validated read across sessions while this process lives."""
    def __init__(self):
        self.lock = threading.Lock()
        self.value = None
        self.attempted = None
        self.stale = False

    def get(self, settings, ttl, force=False, reader=read_source):
        with self.lock:
            now = time.monotonic()
            if force or self.attempted is None or now - self.attempted >= ttl:
                self.attempted = now
                try:
                    value = reader(settings)
                except Exception:
                    self.stale = True
                else:
                    self.value, self.stale = value, False
            return self.value, self.stale


def analytical_xlsx(records, summary, tiers):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados"
    fields = [("name", "Conciliador"), ("qias", "QIAs"), ("changes", "Trocas"), ("credit", "Crédito"),
              ("neo", "Neoenergia"), ("cash", "Caixa"), ("ticket", "Ticket médio"),
              ("qias_projection", "Projeção QIAs"), ("changes_projection", "Projeção trocas"),
              ("monthly_award", "Premiação mensal atual"), ("projected_award", "Mensal projetada"),
              ("weekly_award", "Semanais encerradas"), ("commission_projection", "Comissão projetada")]
    fields += [("cash_projection", "Projeção caixa"), ("tier", "Faixa atual"),
               ("projected_tier", "Faixa projetada"), ("qias_remaining", "QIAs restantes"),
               ("changes_remaining", "Trocas restantes"), ("cash_invalid", "Valores de caixa inválidos")]
    ws.append([label for _, label in fields])
    for row in summary:
        ws.append([float(row[k]) if isinstance(row[k], Decimal) else row[k] for k, _ in fields])
    daily = wb.create_sheet("Diário")
    daily.append(["Data", "Conciliador", "QIAs", "Trocas", "Crédito", "Neoenergia", "Caixa", "NR", "1 A 3", "4 A 6", "Outras", "Valores de caixa inválidos"])
    groups = {}
    for row in records:
        groups.setdefault((row["date"], row["key"]), []).append(row)
    for (day, _), rows in sorted(groups.items()):
        totals = aggregate(rows)
        daily.append([day, rows[0]["name"]] + [float(totals[k]) for k in ("qias", "changes", "credit", "neo", "cash", "NR", "1 A 3", "4 A 6", "OUTRAS", "cash_invalid")])
    issues = wb.create_sheet("Pendências de caixa")
    issues.append(["Linha na aba operacional", "Data", "Conciliador", "Pendência"])
    for row in records:
        if row.get("cash_invalid"):
            issues.append([row["source_line"], row["date"], row["name"], "Valor inválido; caixa e ticket parciais. Corrigir coluna E na origem."])
    config = wb.create_sheet("Faixas consultadas")
    config.append(["Tipo", "Faixa", "QIAs", "Trocas", "Valor"])
    for section, values in tiers.items():
        for i, tier in enumerate(values, 1):
            config.append([section, i, tier["qias"], tier["changes"], float(tier["award"])])
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Untrusted names must never become formulas.
                if isinstance(cell.value, date):
                    cell.number_format = "dd/mm/yyyy"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="065F46")
            sheet.column_dimensions[cell.column_letter].width = 24
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def period_ranking(records, names, start, end, today):
    rows = [{"name": name, "key": key, **aggregate([r for r in records if r["key"] == key and start <= r["date"] <= min(end, today)])} for key, name in names.items()]
    return sorted(rows, key=lambda r: (-r["qias"], -r["changes"], -r["cash"], normalized(r["name"])))


def weekly_ranking(records, names, start, end, today, tiers):
    rows = period_ranking(records, names, start, end, today)
    elapsed, total = useful_days(start, min(today, end)), useful_days(start, end)
    for row in rows:
        row["tier"], row["earned_award"] = award_for(row["qias"], row["changes"], tiers)
        for metric in ("qias", "changes", "cash"):
            row[metric + "_projection"] = Decimal(row[metric]) * total / elapsed if elapsed else Decimal(0)
        row["projected_tier"], projected = award_for(row["qias_projection"], row["changes_projection"], tiers)
        row["award"] = row["earned_award"] if end < today else projected
        row["closed_award"] = row["earned_award"] if end < today else Decimal(0)
    return rows


def prepare_summary(records, tiers, year, month, today, registry):
    rows = summarize(records, tiers, year, month, today, registry)
    for row in rows:
        row["qias_goal"] = tiers["monthly"][0]["qias"]
        row["goal_percent"] = row["qias_projection"] * 100 / row["qias_goal"]
        row["color"] = projection_color(row["qias_projection"], row["qias_goal"])
    return rows


def render_management(st, registry, goals, save_registry, save_goals, password, payload, stale, year, month, summary, month_rows):
    if st.button("VOLTAR AO PAINEL", key="conc_back"):
        st.session_state.conc_management = False
        st.rerun()
    if not st.session_state.get("gestor_autenticado"):
        supplied = st.text_input("Senha do gestor", type="password", key="conc_password")
        if st.button("ENTRAR NA GESTÃO", key="conc_login"):
            import hmac
            if password and hmac.compare_digest(supplied, password):
                st.session_state.gestor_autenticado = True
                st.rerun()
            else:
                st.error("Senha inválida ou não configurada.")
        return
    st.markdown("#### GESTÃO DA CONCILIAÇÃO")
    with st.form("conc_goals"):
        a, b = st.columns(2)
        qias = a.number_input("Meta mensal geral de QIA’s", min_value=0, value=int(goals.get("qias", 0)), step=1, key="conc_goal_qias")
        changes = b.number_input("Meta mensal geral de trocas", min_value=0, value=int(goals.get("changes", 0)), step=1, key="conc_goal_changes")
        st.caption("Metas gerais permanecem salvas até você editá-las. As faixas individuais e premiações continuam na aba Config.")
        if st.form_submit_button("SALVAR METAS GERAIS"):
            try:
                save_goals({"qias": qias, "changes": changes})
            except Exception:
                st.error("Não foi possível salvar as metas. Os valores anteriores foram mantidos.")
            else:
                st.session_state.conc_saved = "Metas gerais salvas."
                st.rerun()
    if st.session_state.get("conc_saved"):
        st.success(st.session_state.pop("conc_saved"))
    if payload is None:
        st.info("A conexão precisa estar disponível para consultar o cadastro e exportar resultados.")
        return
    records, tiers = payload["records"], payload["tiers"]
    st.caption(f"Conexão: {'desatualizada' if stale else 'ativa'} · {len(records)} registros · {len(tiers['monthly'])} faixas mensais · {len(tiers['weekly'])} semanais")
    issues = [r["source_line"] for r in records if r.get("cash_invalid")]
    if issues:
        st.warning("Corrigir valores da coluna E na aba operacional, linhas: " + ", ".join(map(str, issues)))
    all_names = {r["key"]: r["name"] for r in records}
    all_names.update({key: flags.get("name", all_names.get(key, key)) for key, flags in registry.items()})
    with st.form("conc_registry"):
        updated = {}
        for key, name in sorted(all_names.items()):
            a,b,c=st.columns([3,1,1]);a.write(name)
            updated[key] = {"name":name,
                            "active":b.checkbox("Ativo",value=registry.get(key,{}).get("active",True),key="conc_active_"+key),
                            "visible":c.checkbox("Exibir",value=registry.get(key,{}).get("visible",True),key="conc_visible_"+key)}
        added=st.text_input("Adicionar conciliador sem lançamentos (nome oficial)",key="conc_new_name")
        if st.form_submit_button("SALVAR CONCILIADORES"):
            if normalized(added):
                updated.setdefault(normalized(added),{"name":" ".join(added.split()),"active":True,"visible":True})
            try:
                save_registry(updated)
            except Exception:
                st.error("Não foi possível salvar o cadastro. As configurações anteriores foram mantidas.")
            else:
                st.rerun()
    st.download_button("BAIXAR RELATÓRIO ANALÍTICO",analytical_xlsx(month_rows,summary,tiers),
                       file_name=f"conciliacao-{year}-{month:02d}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def render_conciliacao(st, registry, save_registry, manager_password, goals=None, save_goals=None):
    import hashlib
    import json
    goals = goals or {}
    st.markdown(STYLE, unsafe_allow_html=True)
    settings = None
    ttl = 300
    try:
        settings = dict(st.secrets["conciliacao"])
        settings["service_account"] = dict(settings["service_account"])
        ttl = max(60, int(settings.get("cache_seconds", 300)))
        identity = hashlib.sha256(json.dumps(settings,sort_keys=True).encode()).hexdigest()
    except (KeyError,FileNotFoundError,ValueError,TypeError):
        settings = None
        st.error("Configure a seção conciliacao e a conta de serviço nos Secrets do Streamlit.")

    @st.cache_resource(show_spinner=False)
    def source_cache(source_identity):
        return SourceCache()

    @st.fragment(run_every=ttl)
    def body():
        if not st.session_state.get("dashboard_autenticado"):
            return
        view = st.session_state.setdefault("conc_view", "VISÃO GERAL")
        managing = st.session_state.get("conc_management", False)
        force = False
        today = datetime.now(TZ).date()
        with st.container(key="dashboard_view_controls"):
            nav, refresh = st.columns([5,1],vertical_alignment="center")
            with nav:
                if not managing:
                    with st.container(key="top_nav_buttons"):
                        for col,label in zip(st.columns(3,gap="small"),("VISÃO GERAL","DIÁRIO","SEMANAL")):
                            if col.button(label,key="conc_nav_"+label,use_container_width=True,type="primary" if view==label else "secondary"):
                                st.session_state.conc_view=label
                                st.rerun()
            with refresh:
                force=st.button("↻ Atualizar dados",key="conc_refresh",help="Consultar novamente os resultados e as faixas da planilha")
        payload,stale=source_cache(identity).get(settings,ttl,force) if settings else (None,True)
        if stale:
            st.warning("Não foi possível atualizar. Última leitura válida mantida, quando disponível.")
        records=payload["records"] if payload else []
        months=sorted({(r["date"].year,r["date"].month) for r in records}|{(today.year,today.month)},reverse=True)
        a,b=st.columns([1,3],vertical_alignment="center")
        with a:
            year,month=st.selectbox("Competência",months,format_func=lambda p:f"{p[1]:02d}/{p[0]}",key="conc_month",label_visibility="collapsed")
        if payload:
            b.caption(f"Atualizado {payload['synced_at']:%d/%m às %H:%M} · automático a cada {ttl//60} min")
        summary=prepare_summary(records,payload["tiers"],year,month,today,registry) if payload else []
        names={r["key"]:r["name"] for r in summary}
        month_rows=[r for r in records if r["key"] in names and (r["date"].year,r["date"].month)==(year,month) and r["date"]<=today]
        if managing:
            render_management(st,registry,goals,save_registry,save_goals,manager_password,payload,stale,year,month,summary,month_rows)
            return
        if payload is None:
            st.error("Ainda não há leitura válida da planilha.")
            return
        if payload["duplicates"]:
            st.caption(f"{payload['duplicates']} registros duplicados desconsiderados.")
        periods=weeks(year,month)
        default_day=today.day if (year,month)==(today.year,today.month) else 1
        if view=="SEMANAL":
            week_key=f"conc_week_{year}_{month}"
            index=st.session_state.setdefault(week_key,min((default_day-1)//7,len(periods)-1))
            with st.container(key="week_nav_buttons"):
                for i,col in enumerate(st.columns(len(periods),gap="small")):
                    if col.button(f"S{i+1}",key=f"conc_week_btn_{year}_{month}_{i}",type="primary" if i==index else "secondary",help=f"{periods[i][0]:%d/%m} a {periods[i][1]:%d/%m}"):
                        st.session_state[week_key]=i
                        st.rerun()
            start,end=periods[index]
            ranked=weekly_ranking(records,names,start,end,today,payload["tiers"]["weekly"])
            for row in ranked:
                row["color"]=projection_color(row["qias_projection"],payload["tiers"]["weekly"][0]["qias"])
            period=f"S{index+1} · {start:%d/%m/%Y} a {end:%d/%m/%Y}"
            st.caption(period + (" · premiação conquistada" if end<today else " · premiação projetada; semana em andamento" if start<=today else " · semana futura"))
        elif view=="DIÁRIO":
            start=end=st.date_input("Dia",date(year,month,default_day),min_value=date(year,month,1),max_value=date(year,month,monthrange(year,month)[1]),key=f"conc_day_{year}_{month}")
            ranked=period_ranking(records,names,start,end,today)
            week={r["key"]:r for r in period_ranking(records,names,*periods[(start.day-1)//7],today)}
            monthly={r["key"]:r for r in summary}
            for row in ranked:
                row.update(color=daily_color(row["qias"]),week_qias=week[row["key"]]["qias"],month_qias=monthly[row["key"]]["qias"])
            period=f"{start:%d/%m/%Y}"
        else:
            start,end=date(year,month,1),date(year,month,monthrange(year,month)[1])
            ranked=summary
            period=f"{month:02d}/{year}"
        filtered=[r for r in month_rows if start<=r["date"]<=end]
        totals=aggregate(filtered)
        if totals["cash_invalid"]:
            st.warning("Caixa e Ticket Médio parciais: há valores inválidos na coluna E. QIAs e trocas foram preservados. Consulte Gestão.")
        if view=="VISÃO GERAL":
            st.markdown(summary_html(totals,summary,goals),unsafe_allow_html=True)
        else:
            from app_core import cards
            period_cards=[("TOTAL DE QIAs HOJE" if start==today and view=="DIÁRIO" else "TOTAL DE QIAs",integer(totals["qias"])),("TOTAL DE TROCAS HOJE" if start==today and view=="DIÁRIO" else "TOTAL DE TROCAS",integer(totals["changes"])),("CRÉDITO / NEO",f"{totals['credit']} / {totals['neo']}"),("Ticket Médio",money(totals["ticket"]))]
            cards(st,period_cards[:2] if view=="DIÁRIO" else period_cards)
        st.markdown(regimes_html(totals),unsafe_allow_html=True)
        if view=="DIÁRIO":
            st.caption("🔴 0–14 · 🟠 15–19 · 🟡 20–24 · 🟢 25–29 · 🔵 30 ou mais QIAs")
            png=ranking_png(ranked,"RANKING DIÁRIO · CONCILIAÇÃO",period,payload["synced_at"],view,totals,goals)
            st.image(png,use_container_width=True)
        else:
            st.markdown(ranking_html(ranked,view),unsafe_allow_html=True)
            png=None
        if view=="SEMANAL" and index:
            previous=aggregate([r for r in month_rows if periods[index-1][0]<=r["date"]<=periods[index-1][1]])
            st.caption(f"Variação para S{index}: {totals['qias']-previous['qias']:+d} QIAs · {totals['changes']-previous['changes']:+d} trocas")
        if view in ("VISÃO GERAL","DIÁRIO","SEMANAL"):
            if png is None:
                png=ranking_png(ranked,"RANKING GERAL · CONCILIAÇÃO" if view=="VISÃO GERAL" else "RANKING SEMANAL · CONCILIAÇÃO",period,payload["synced_at"],view,totals,goals)
            st.download_button("BAIXAR RANKING PNG",png,file_name=f"ranking-conciliacao-{view.lower()}-{start.isoformat()}.png",mime="image/png",key="conc_download")
    body()
